from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: verify_xlsx.py <workbook.xlsx>")
    source = Path(sys.argv[1]).resolve()
    if not source.exists():
        raise SystemExit(f"workbook not found: {source}")
    workbook = load_workbook(source, data_only=False)
    expected = ["Detected table", "Raw OCR", "Unassigned OCR", "Diagnostics"]
    if workbook.sheetnames != expected:
        raise AssertionError(f"unexpected sheets: {workbook.sheetnames}")
    table = workbook["Detected table"]
    if table["A1"].value != "Merged heading" or table["C2"].value != "125.00":
        raise AssertionError("key cell values were not preserved")
    if "A1:B1" not in {str(item) for item in table.merged_cells.ranges}:
        raise AssertionError("merged cell was not preserved")
    if table.max_row < 2 or table.max_column < 3:
        raise AssertionError("blank grid coordinates were compacted")
    result = {"openpyxl": "passed", "sheets": workbook.sheetnames, "merge": "A1:B1", "blankCoordinate": "B2", "libreOffice": "not_available"}
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable:
        with tempfile.TemporaryDirectory() as directory:
            completed = subprocess.run([executable, "--headless", "--convert-to", "xlsx", "--outdir", directory, str(source)], capture_output=True, text=True, timeout=90)
            combined = f"{completed.stdout}\n{completed.stderr}".lower(); produced = Path(directory, source.name).exists()
            if completed.returncode != 0 or (not produced and "exists" not in combined and "source format" not in combined):
                raise AssertionError(f"LibreOffice could not parse workbook: {completed.stdout} {completed.stderr}")
            result["libreOffice"] = "passed"
    output = source.with_suffix(".compatibility.json"); output.write_text(json.dumps(result, indent=2), encoding="utf-8"); print(json.dumps(result, indent=2))


if __name__ == "__main__": main()
